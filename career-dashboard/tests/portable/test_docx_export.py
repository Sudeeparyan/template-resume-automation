"""Word copies of a tailored CV and cover letter (services/docx_export.py): the same content as the
checked PDF's source, as plain paragraphs and real bullets that applicant-tracking systems read."""

from __future__ import annotations

import hashlib
import io
import json

import pytest
import yaml
from docx import Document

from backend.services import docx_export
from backend.services.intake import resume_base
from backend.services.resume_studio import ResumeStudio
from career import tex_escape
from test_hunt import JD, ireland_profile

ROLE_FACTS = [
    "Built 12 Power BI dashboards used by 40 store managers across Leinster",
    "Cut weekly reporting time by 30% with scheduled SQL automation",
    "Reconciled 2 million sales rows each month & flagged #duplicate orders for finance",
]
PROJECT_BULLETS = [
    "Forecast emergency department waits with an MAE of 7.3 minutes",
    "Compared ARIMA & gradient boosting on three years of hourly data",
    "Presented the findings to two clinical leads at a Dublin hospital",
]
CANDIDATE = {"full_name": "Aoife Ní Bhriain", "email": "aoife_nb@example.com", "phone": "+353 1 555 0100",
             "linkedin": "https://www.linkedin.com/in/aoife-example"}


def evidence() -> dict:
    def claim(id, category, **fields):
        return {"id": id, "status": "user_reported", "category": category, "approved_external_use": "always",
                "source_refs": ["me/resume.pdf"], **fields}

    return {
        "candidate_revision": "docx-test-1",
        "claims": [
            claim("IDENTITY-001", "identity", value=CANDIDATE["full_name"]),
            claim("CONTACT-EMAIL-001", "contact", value=CANDIDATE["email"]),
            claim("CONTACT-PHONE-001", "contact", value=CANDIDATE["phone"]),
            claim("CONTACT-LINKEDIN-001", "contact", value=CANDIDATE["linkedin"]),
            claim("EDU-MSC-001", "education", institution="University College Dublin", location="Dublin, Ireland",
                  degree_as_supplied="MSc Business Analytics", dates="Sep 2025 - Aug 2026"),
            claim("COURSE-001", "academic_coursework", degree_id="EDU-MSC-001",
                  approved_facts=["Data Mining", "Optimisation", "Big Data Systems"]),
            claim("SKILL-001", "skill", title="Analytics & BI", approved_facts=["SQL", "Power BI", "Python", "Excel"]),
            claim("SKILL-002", "skill", title="Data Engineering", approved_facts=["dbt", "Airflow", "PostgreSQL"]),
            claim("EMP-001", "employment", title="Data Analyst", employer="Example Retail Pvt Ltd", location="Pune, India",
                  dates="Jan 2022 - Jul 2025", approved_facts=ROLE_FACTS),
        ],
        "projects": [
            {"id": "PROJ-001", "status": "user_reported", "title": "Hospital wait-time forecast",
             "resume_content": {"title": "Hospital Wait-Time Forecast", "context": "MSc project, 2026",
                                "bullets": PROJECT_BULLETS}},
            {"id": "PROJ-002", "status": "user_reported", "title": "Dublin Bikes demand",
             "resume_content": {"title": "Dublin Bikes Demand", "context": "Personal project",
                                "bullets": ["Modelled hourly demand at 110 docking stations with LightGBM",
                                            "Served the predictions through a small FastAPI endpoint"]}},
        ],
    }


def profile() -> dict:
    return {"candidate": CANDIDATE, "target_roles": {"primary": ["Data Analyst"]},
            "resume_contract": {"paper": "a4", "header_fields": ["full_name", "email", "phone", "linkedin"]}}


def body_text(data: bytes) -> list[tuple[str, str]]:
    """(style, text) for every paragraph of a .docx."""
    return [(paragraph.style.name, paragraph.text) for paragraph in Document(io.BytesIO(data)).paragraphs]


def test_plain_text_resolves_escapes_accents_and_layout_commands():
    assert docx_export.plain(r"R\&D \textbar{} 30\% faster \textbackslash{} C\# \$5") == "R&D | 30% faster \\ C# $5"
    assert docx_export.plain(r"\vspace{3pt}Caf\'{e} \v{S}koda \c{c}a\hspace*{1em}") == "Café Škoda ça"
    assert docx_export.plain(r"\href{mailto:a\_b@example.com}{a\_b@example.com}") == "a_b@example.com"
    assert docx_export.plain(r"Jan 2024 -- Present \textasciitilde{}1 \textasciicircum{}2 \{x\}") == "Jan 2024 – Present ~1 ^2 {x}"
    assert docx_export.plain(r"R\textsuperscript{2} of 0.81 \textit{(test set)}") == "R² of 0.81 (test set)"
    assert docx_export.plain(r"\color{gray}\setlength{\parskip}{2pt}\textbf{Kept}") == "Kept"


def test_a_rendered_resume_becomes_the_same_sections_roles_bullets_and_skills():
    source = resume_base.render(profile(), evidence())
    found = docx_export.blocks(source)
    kinds = [block["type"] for block in found]
    assert kinds[:2] == ["name", "contact"] and found[0]["text"] == CANDIDATE["full_name"]
    assert found[1]["text"] == " | ".join([CANDIDATE["email"], CANDIDATE["phone"], CANDIDATE["linkedin"]])
    assert [block["text"] for block in found if block["type"] == "section"] == [
        "Education", "Technical Skills", "Professional Experience", "Projects"]
    roles = [block for block in found if block["type"] == "role"]
    assert roles[0] == {"type": "role", "title": "University College Dublin", "dates": "Sep 2025 – Aug 2026",
                        "org": "MSc Business Analytics", "place": "Dublin, Ireland"}
    assert roles[1]["org"] == "Example Retail Pvt Ltd" and roles[1]["dates"] == "Jan 2022 – Jul 2025"
    labelled = {block["label"]: block["text"] for block in found if block["type"] == "labelled"}
    assert labelled == {"Coursework:": "Data Mining, Optimisation, Big Data Systems",
                        "Analytics & BI:": "SQL, Power BI, Python, Excel",
                        "Data Engineering:": "dbt, Airflow, PostgreSQL"}
    headings = [(block["title"], block["right"]) for block in found if block["type"] == "heading"]
    assert headings == [("Hospital Wait-Time Forecast", "MSc project, 2026"), ("Dublin Bikes Demand", "Personal project")]
    bullets = [block["text"] for block in found if block["type"] == "bullet"]
    # Every \item in the PDF's source is one bullet here, word for word (escapes resolved).
    assert len(bullets) == source.split("\\begin{document}")[1].count("\\item ")
    printed = [fact for fact in ROLE_FACTS + PROJECT_BULLETS if tex_escape(fact) in source]
    assert len(printed) >= 4 and all(fact in bullets for fact in printed)
    assert not any("\\" in str(value) or "{" in str(value) for block in found for value in block.values())


def test_the_word_cv_is_a_single_column_a4_document_with_real_bullets():
    source = resume_base.render(profile(), evidence())
    data = docx_export.cv_document(source, title="CV for Data Analyst at Acme")
    document = Document(io.BytesIO(data))
    assert document.core_properties.title == "CV for Data Analyst at Acme"
    assert document.core_properties.author == CANDIDATE["full_name"]
    section = document.sections[0]
    assert (round(section.page_width.mm), round(section.page_height.mm)) == (210, 297)
    assert not document.tables and not document.inline_shapes
    paragraphs = body_text(data)
    bullets = [text for style, text in paragraphs if style == "List Bullet"]
    assert bullets and all(fact in bullets for fact in PROJECT_BULLETS[:2])
    assert ("Normal", "Data Analyst\tJan 2022 – Jul 2025") in paragraphs
    assert "Hospital Wait-Time Forecast\tMSc project, 2026" in docx_export.document_text(data)


def test_a_cover_letter_keeps_its_paragraphs_and_adds_nothing():
    letter = "Dear Hiring Manager,\n\nI am applying for the Data Analyst role.\nLine two of the same paragraph.\n\n\nKind regards,\nAoife"
    data = docx_export.letter_document(letter, title="Cover letter", author="Aoife")
    paragraphs = [text for _, text in body_text(data)]
    assert paragraphs == ["Dear Hiring Manager,", "I am applying for the Data Analyst role.\nLine two of the same paragraph.",
                          "Kind regards,\nAoife"]


def compiled_draft(services, job, source):
    """A saved Studio draft whose revision 2 compiled: what the dashboard holds after Build."""
    w = services.w
    ResumeStudio(services)
    folder = w.root / "data/output/applications" / job["id"] / "studio"
    preview = folder / "preview-2"
    preview.mkdir(parents=True)
    (preview / "resume.pdf").write_bytes(b"%PDF-1.4 compiled fixture")
    (folder / "preview.json").write_text(json.dumps({
        "revision": 2, "source_sha256": hashlib.sha256(source.encode()).hexdigest(), "page_count": 1,
        "path": preview.relative_to(w.root / "data/output").as_posix(), "layout": {"full_pages": True}}), encoding="utf-8")
    with w.connect() as db:
        db.execute("UPDATE jobs SET folder=? WHERE id=?", (folder.parent.relative_to(w.root).as_posix(), job["id"]))
        db.execute("INSERT INTO studio_drafts VALUES(?,?,?,?,?,?)",
                   (job["id"], source, 2, folder.relative_to(w.root).as_posix(), services.now(), "docx-test-1"))
    return preview


def test_the_studio_downloads_a_word_copy_of_the_current_compiled_revision_only(tmp_path, monkeypatch):
    services = ireland_profile(tmp_path)
    (services.w.root / "data/context/evidence.yml").write_text(yaml.safe_dump(evidence(), allow_unicode=True), encoding="utf-8")
    job = services.w.add_job("Acme Analytics", "Data Analyst", "Dublin, Ireland", "https://jobs.example/acme/1", JD)
    source = resume_base.render(profile(), evidence())
    preview = compiled_draft(services, job, source)
    studio = ResumeStudio(services)
    path, name = studio.download(job["id"], "docx")
    assert path == preview / "resume.docx" and name == "acme-analytics-data-analyst-resume-v2.docx"
    text = docx_export.document_text(path.read_bytes())
    assert CANDIDATE["full_name"] in text and PROJECT_BULLETS[0] in text and "\\" not in text
    # A draft edited after its last build has no checked PDF, so no Word copy either.
    with services.w.connect() as db:
        db.execute("UPDATE studio_drafts SET source=? WHERE job_id=?", (source + "\n% edited", job["id"]))
    with pytest.raises(ValueError, match="current successfully compiled revision"):
        studio.download(job["id"], "docx")
    with services.w.connect() as db:
        db.execute("UPDATE studio_drafts SET source=? WHERE job_id=?", (source, job["id"]))
    monkeypatch.setenv("CAREER_FEATURES", "-docx_export")
    with pytest.raises(ValueError, match="switched off"):
        studio.download(job["id"], "docx")
