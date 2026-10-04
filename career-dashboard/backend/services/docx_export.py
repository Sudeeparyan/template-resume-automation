"""Word (.docx) copies of a tailored CV and cover letter, with the same content as the PDF.

Many Irish employers and agencies ask for a Word CV. The DOCX is made from the same draft
source as the PDF (the Resume Studio's LaTeX), so it never says anything the checked PDF does
not: the same sections, roles, dates, bullets and skills, in the same order. It is a plain,
single-column A4 document that applicant-tracking systems read reliably: standard headings,
real bullet lists, no tables, text boxes, columns or images.

Only the constructs the resume templates write are understood (``\\section``, ``\\roleheading``,
``\\clientheading``, bold-label skill lines, project titles, ``resumeitems`` lists and the
centred name and contact lines); anything else is kept as plain text, never dropped silently.
"""

from __future__ import annotations

import io
import re
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

FONT = "Calibri"
BODY_PT = 10.5
_ACCENTS = {"'": "\u0301", "`": "\u0300", "^": "\u0302", '"': "\u0308", "~": "\u0303", "c": "\u0327", "v": "\u030c"}
_SYMBOLS = {r"\textbullet{}": "\u2022", r"\textperiodcentered{}": "\u00b7", r"\textdegree{}": "\u00b0",
            r"$\times$": "\u00d7", r"$\rightarrow$": "\u2192", r"$\le$": "\u2264", r"$\ge$": "\u2265",
            r"\texteuro{}": "\u20ac", r"\pounds{}": "\u00a3", r"$|$": "|", r"\textbar{}": "|", r"\textbar": "|"}
# Characters the resume writer escapes (career.tex_escape). They are held as private-use
# characters while commands and grouping braces are removed, then put back.
_LITERALS = {r"\textbackslash{}": "\\", r"\textasciitilde{}": "~", r"\textasciicircum{}": "^"}
_HOLD = {character: chr(0xE000 + number) for number, character in enumerate("%&_#${}\\~^")}
_RESTORE = {ord(held): character for character, held in _HOLD.items()}
# Layout commands whose arguments are measurements or settings, never text: (arguments to drop).
_LAYOUT = {"vspace": 1, "hspace": 1, "addvspace": 1, "enlargethispage": 1, "label": 1, "color": 1, "textcolor": 1,
           "pagestyle": 1, "thispagestyle": 1, "linespread": 1, "phantom": 1, "hphantom": 1, "vphantom": 1,
           "parbox": 1, "setlength": 2, "addtolength": 2, "settowidth": 2, "setcounter": 2, "fontsize": 2,
           "rule": 2, "titlespacing": 4}
_LAYOUT_COMMAND = re.compile(r"\\(" + "|".join(sorted(_LAYOUT, key=len, reverse=True)) + r")\*?(?![A-Za-z@])(?:\s*\[[^\]]*\])*")


def _argument(text: str, start: int) -> tuple[str, int]:
    """The brace group starting at ``start`` (a "{") and the index after it."""
    depth, index = 0, start
    while index < len(text):
        character = text[index]
        if character == "\\":
            index += 2
            continue
        if character == "{":
            depth += 1
        elif character == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1:index], index + 1
        index += 1
    return text[start + 1:], len(text)


def _arguments(text: str, start: int, count: int) -> tuple[list[str], int]:
    values, index = [], start
    for _ in range(count):
        while index < len(text) and text[index].isspace():
            index += 1
        if index >= len(text) or text[index] != "{":
            break
        value, index = _argument(text, index)
        values.append(value)
    return values, index


def _drop_layout(text: str) -> str:
    """Remove layout commands with their measurement arguments (``\\vspace{3pt}`` is not text)."""
    out, index = [], 0
    while match := _LAYOUT_COMMAND.search(text, index):
        out.append(text[index:match.start()])
        _, index = _arguments(text, match.end(), _LAYOUT[match[1]])
    return "".join(out) + text[index:]


def _accent(match) -> str:
    return unicodedata.normalize("NFC", match[2] + _ACCENTS[match[1]])


def plain(latex: str) -> str:
    """Display text of a LaTeX fragment: escapes, accents, dashes and links resolved, commands removed."""
    text = re.sub(r"\\\\(?:\[[^\]]*\])?", "\n", latex)  # line breaks first: "\\" never starts an escape
    for command, character in _LITERALS.items():
        text = text.replace(command, _HOLD[character])
    for escaped in "%&_#${}":
        text = text.replace("\\" + escaped, _HOLD[escaped])
    for symbol, character in _SYMBOLS.items():
        text = text.replace(symbol, character)
    text = _drop_layout(text)
    text = re.sub(r"\\href\{([^{}]*)\}\{([^{}]*)\}", r"\2", text)
    text = re.sub(r"\\url\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\textsuperscript\{([^{}]*)\}", lambda m: {"2": "\u00b2", "3": "\u00b3"}.get(m[1], m[1]), text)
    text = re.sub(r"\\(['`^\"~])\{?([A-Za-z])\}?", _accent, text)
    text = re.sub(r"\\([cv])\{([A-Za-z])\}", _accent, text)  # \c{c}, \v{s}; never \color or \vspace
    text = text.replace("---", "\u2014").replace("--", "\u2013").replace("~", "\u00a0").replace(r"\,", "\u2009")
    text = re.sub(r"\\(?:hfill|quad|qquad|newline)(?![A-Za-z@])", "  ", text)
    text = re.sub(r"\\(?:begin|end)\{[^{}]*\}", "", text)
    text = re.sub(r"\\[A-Za-z@]+\*?(?:\[[^\]]*\])*", "", text)
    text = text.replace("{", "").replace("}", "").translate(_RESTORE)
    text = text.replace("``", "\u201c").replace("''", "\u201d")
    return "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")).strip()


def _strip_comments(source: str) -> str:
    from validate_resume import strip_latex_comments

    return strip_latex_comments(source)


def _expand(source: str, body: str) -> str:
    from validate_resume import extract_zero_argument_macros

    macros = extract_zero_argument_macros(source)
    for _ in range(4):  # a macro may name another
        before = body
        for name, value in sorted(macros.items(), key=lambda item: len(item[0]), reverse=True):
            body = re.sub(rf"\\{re.escape(name)}(?![A-Za-z@])", lambda _m, v=value: v, body)
        if body == before:
            break
    return body


def blocks(source: str) -> list[dict]:
    """The document body as a list of typed blocks, in order."""
    text = _strip_comments(source)
    body = text.split("\\begin{document}", 1)[-1].split("\\end{document}", 1)[0]
    body = _expand(text, body)
    out: list[dict] = []
    centre = re.search(r"\\begin\{center\}(.*?)\\end\{center\}", body, re.S)
    if centre:
        lines = [plain(part) for part in re.split(r"\\\\(?:\[[^\]]*\])?", centre.group(1))]
        lines = [line for line in lines if line]
        if lines:
            out.append({"type": "name", "text": lines[0]})
        out += [{"type": "contact", "text": line} for line in lines[1:]]
        body = body[:centre.start()] + body[centre.end():]
    index, in_list = 0, False
    pending = ""

    def flush():
        nonlocal pending
        words = plain(pending)
        if words:
            out.append({"type": "text", "text": words})
        pending = ""

    while index < len(body):
        if body.startswith("\\section", index):
            flush()
            (title,), index = _arguments(body, index + len("\\section") + (1 if body[index + 8:index + 9] == "*" else 0), 1)
            out.append({"type": "section", "text": plain(title)})
            continue
        if body.startswith("\\roleheading", index):
            flush()
            values, index = _arguments(body, index + len("\\roleheading"), 4)
            values += [""] * (4 - len(values))
            out.append({"type": "role", "title": plain(values[0]), "dates": plain(values[1]),
                        "org": plain(values[2]), "place": plain(values[3])})
            continue
        if body.startswith("\\clientheading", index):
            flush()
            (client,), index = _arguments(body, index + len("\\clientheading"), 1)
            out.append({"type": "client", "text": plain(client)})
            continue
        if body.startswith("\\begin{resumeitems}", index) or body.startswith("\\begin{itemize}", index):
            flush()
            in_list = True
            index = body.index("}", index) + 1
            options = re.match(r"\s*\[[^\]]*\]", body[index:])  # \begin{itemize}[leftmargin=*]
            index += options.end() if options else 0
            continue
        if body.startswith("\\end{resumeitems}", index) or body.startswith("\\end{itemize}", index):
            if in_list and pending.strip():
                out.append({"type": "bullet", "text": plain(pending)})
                pending = ""
            in_list = False
            index = body.index("}", index) + 1
            continue
        if in_list and body.startswith("\\item", index) and not body[index + 5:index + 6].isalpha():
            if pending.strip():
                out.append({"type": "bullet", "text": plain(pending)})
            pending = ""
            index += len("\\item")
            continue
        if not in_list and body.startswith("\\textbf", index):
            flush()
            (label,), after = _arguments(body, index + len("\\textbf"), 1)
            end = body.find("\n\n", after)
            line_end = body.find("\n", after)
            stop = min(x for x in (end, line_end, len(body)) if x >= 0)
            rest = body[after:stop]
            if "\\hfill" in rest:  # a project title line: title, then its context on the right
                left, _, right = rest.partition("\\hfill")
                out.append({"type": "heading", "title": plain(label + left), "right": plain(right)})
            else:
                out.append({"type": "labelled", "label": plain(label), "text": plain(rest)})
            index = stop
            continue
        if body[index] == "\n" and not in_list and body[index:index + 2] == "\n\n":
            flush()
        pending += body[index]
        index += 1
    flush()
    return [block for block in out if any(str(value).strip() for key, value in block.items() if key != "type")]


# ---- writing ------------------------------------------------------------------------------------

def _base_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(17)
    section.top_margin = section.bottom_margin = Mm(15)
    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY_PT)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.space_before = Pt(0)
    return document


def _bottom_rule(paragraph) -> None:
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _right_tab(document, paragraph) -> None:
    section = document.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)


def _left_right(document, left: str, right: str, *, bold=False, italic=False):
    paragraph = document.add_paragraph()
    _right_tab(document, paragraph)
    run = paragraph.add_run(left)
    run.bold, run.italic = bold, italic
    if right:
        tail = paragraph.add_run("\t" + right)
        tail.italic = italic
    paragraph.paragraph_format.space_after = Pt(0)
    return paragraph


def cv_document(source: str, *, title: str = "CV") -> bytes:
    """The .docx bytes of one resume draft (its LaTeX source)."""
    document = _base_document()
    document.core_properties.title = title
    for block in blocks(source):
        kind = block["type"]
        if kind == "name":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(block["text"])
            run.bold, run.font.size = True, Pt(18)
            document.core_properties.author = block["text"]
        elif kind == "contact":
            paragraph = document.add_paragraph(block["text"])
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "section":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(8)
            run = paragraph.add_run(block["text"])
            run.bold, run.font.size, run.font.small_caps = True, Pt(12), True
            _bottom_rule(paragraph)
        elif kind == "role":
            _left_right(document, block["title"], block["dates"], bold=True).paragraph_format.space_before = Pt(4)
            if block["org"] or block["place"]:
                _left_right(document, block["org"], block["place"], italic=True)
        elif kind == "client":
            document.add_paragraph().add_run(block["text"]).bold = True
        elif kind == "heading":
            _left_right(document, block["title"], block["right"], bold=True).paragraph_format.space_before = Pt(4)
        elif kind == "labelled":
            paragraph = document.add_paragraph()
            paragraph.add_run(block["label"] + " ").bold = True
            paragraph.add_run(block["text"])
        elif kind == "bullet":
            document.add_paragraph(block["text"], style="List Bullet")
        else:
            document.add_paragraph(block["text"])
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def letter_document(text: str, *, title: str = "Cover letter", author: str = "") -> bytes:
    """A cover letter as .docx: its paragraphs as written, nothing added."""
    document = _base_document()
    document.core_properties.title = title
    if author:
        document.core_properties.author = author
    for paragraph in re.split(r"\n\s*\n", text.strip()):
        lines = [line.rstrip() for line in paragraph.splitlines() if line.strip()]
        if not lines:
            continue
        current = document.add_paragraph()
        current.paragraph_format.space_after = Pt(8)
        for number, line in enumerate(lines):
            if number:
                current.add_run().add_break()
            current.add_run(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def document_text(data: bytes) -> str:
    """All paragraph text of a .docx (used to check a copy against its PDF)."""
    document = Document(io.BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)
